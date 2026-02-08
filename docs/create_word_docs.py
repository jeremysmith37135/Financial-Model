"""
Create Word documents from the markdown documentation
"""
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import os

def create_user_manual():
    """Create the User Manual Word document"""
    doc = Document()

    # Title
    title = doc.add_heading('CFO Financial Model Generator', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph('User Manual v3.2.13')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    # Table of Contents
    doc.add_heading('Table of Contents', level=1)
    toc_items = [
        '1. Introduction',
        '2. System Requirements',
        '3. Getting Started',
        '4. Main Interface',
        '5. Single Entity Mode',
        '6. Multi-Division Mode',
        '7. Update Existing Model',
        '8. Generated Excel Workbook',
        '9. Working with the Output',
        '10. Troubleshooting',
        '11. FAQ'
    ]
    for item in toc_items:
        doc.add_paragraph(item, style='List Number')

    doc.add_page_break()

    # Section 1: Introduction
    doc.add_heading('1. Introduction', level=1)

    doc.add_heading('What is the CFO Financial Model Generator?', level=2)
    doc.add_paragraph(
        'The CFO Financial Model Generator is a desktop application designed for CFOs, '
        'controllers, and financial professionals to create comprehensive financial models '
        'from QuickBooks or other accounting system exports.'
    )

    doc.add_paragraph('The application transforms your P&L and Balance Sheet data into a fully-functional Excel workbook with:')
    features = [
        'Interactive P&L and Balance Sheet reports with YTD calculations',
        'Cash Flow Statement automatically generated from P&L and Balance Sheet data',
        'Executive Dashboard with key financial metrics and charts',
        'Forecast module comparing Actuals vs Budget with variance analysis',
        'Multi-division consolidation for companies with multiple entities'
    ]
    for feature in features:
        doc.add_paragraph(feature, style='List Bullet')

    doc.add_heading('Key Features', level=2)

    # Features table
    table = doc.add_table(rows=7, cols=2)
    table.style = 'Table Grid'

    headers = table.rows[0].cells
    headers[0].text = 'Feature'
    headers[1].text = 'Description'

    features_data = [
        ('Auto-Detection', 'Automatically detects date ranges and account structures'),
        ('File Validation', 'Validates that P&L and Balance Sheet files are correctly assigned'),
        ('Dynamic Formulas', 'All reports use formulas - change source data and reports update'),
        ('Period Selection', 'Select any period from the dropdown to view historical data'),
        ('Multi-Division', 'Consolidate multiple divisions/departments into one model'),
        ('Update Capability', 'Add new months to existing models without rebuilding')
    ]

    for i, (feature, desc) in enumerate(features_data, 1):
        row = table.rows[i].cells
        row[0].text = feature
        row[1].text = desc

    doc.add_page_break()

    # Section 2: System Requirements
    doc.add_heading('2. System Requirements', level=1)

    doc.add_heading('Minimum Requirements', level=2)
    requirements = [
        'Operating System: Windows 10 or Windows 11',
        'Microsoft Excel: Microsoft Excel 2016 or later (required)',
        'Memory: 4 GB RAM minimum (8 GB recommended)',
        'Disk Space: 200 MB for application + space for generated files',
        'Screen Resolution: 1280 x 720 minimum'
    ]
    for req in requirements:
        doc.add_paragraph(req, style='List Bullet')

    doc.add_heading('Important Notes', level=2)
    notes = [
        'Excel must be installed - The application uses Excel\'s COM interface',
        'Close Excel before running - For best results, close any open workbooks',
        'Enable macros - Generated workbooks contain VBA macros for navigation'
    ]
    for note in notes:
        doc.add_paragraph(note, style='List Bullet')

    doc.add_page_break()

    # Section 3: Getting Started
    doc.add_heading('3. Getting Started', level=1)

    doc.add_heading('Installation', level=2)
    steps = [
        'Download CFO_Financial_Model_Generator_v3.2.13.exe',
        'Place the EXE file in a permanent location (e.g., C:\\Program Files\\CFO Tools\\)',
        'Create a desktop shortcut if desired',
        'Double-click to run - no installation required'
    ]
    for i, step in enumerate(steps, 1):
        doc.add_paragraph(f'{i}. {step}')

    doc.add_heading('Preparing Your Data', level=2)

    doc.add_heading('P&L Export', level=3)
    doc.add_paragraph('Export your P&L from QuickBooks or your accounting system with:')
    pl_reqs = [
        'Columns: Account names in column A, months as column headers',
        'Rows: All income and expense accounts',
        'Format: Excel file (.xlsx or .xls)'
    ]
    for req in pl_reqs:
        doc.add_paragraph(req, style='List Bullet')

    doc.add_heading('Balance Sheet Export', level=3)
    doc.add_paragraph('Export your Balance Sheet with:')
    bs_reqs = [
        'Columns: Account names in column A, months as column headers',
        'Rows: All asset, liability, and equity accounts',
        'Format: Excel file (.xlsx or .xls)'
    ]
    for req in bs_reqs:
        doc.add_paragraph(req, style='List Bullet')

    doc.add_page_break()

    # Section 4: Main Interface
    doc.add_heading('4. Main Interface', level=1)
    doc.add_paragraph(
        'When you launch the application, you\'ll see the main window with the following elements:'
    )

    interface_elements = [
        ('Mode', 'Shows "Single Entity" or "X Divisions configured"'),
        ('Configure Divisions', 'Opens the multi-division setup dialog'),
        ('Company Name', 'Enter the company name for report headers'),
        ('Model File', '(Optional) Select existing model to update'),
        ('P&L File', 'Select your P&L export file'),
        ('Balance Sheet', 'Select your Balance Sheet export file'),
        ('Generate', 'Creates the financial model'),
        ('Progress Bar', 'Shows generation progress with time estimate'),
        ('Status', 'Displays current operation or errors')
    ]

    table = doc.add_table(rows=len(interface_elements)+1, cols=2)
    table.style = 'Table Grid'
    headers = table.rows[0].cells
    headers[0].text = 'Element'
    headers[1].text = 'Description'

    for i, (element, desc) in enumerate(interface_elements, 1):
        row = table.rows[i].cells
        row[0].text = element
        row[1].text = desc

    doc.add_page_break()

    # Section 5: Single Entity Mode
    doc.add_heading('5. Single Entity Mode', level=1)
    doc.add_paragraph(
        'Single Entity mode is for companies with one set of books - no divisions, '
        'departments, or branches to consolidate.'
    )

    doc.add_heading('Step-by-Step Guide', level=2)

    doc.add_paragraph('Step 1: Enter Company Name', style='Heading 3')
    doc.add_paragraph('Type your company name in the "Company Name" field. This appears on all report headers.')

    doc.add_paragraph('Step 2: Select P&L File', style='Heading 3')
    doc.add_paragraph('Click Browse next to "P&L File", navigate to your P&L export, and select it.')

    doc.add_paragraph('Step 3: Select Balance Sheet File', style='Heading 3')
    doc.add_paragraph('Click Browse next to "Balance Sheet", navigate to your BS export, and select it.')

    doc.add_paragraph('Step 4: Generate Model', style='Heading 3')
    doc.add_paragraph('Click "Generate Financial Model", choose save location, and wait for completion.')

    doc.add_page_break()

    # Section 6: Multi-Division Mode
    doc.add_heading('6. Multi-Division Mode', level=1)
    doc.add_paragraph(
        'Multi-Division mode is for companies with multiple entities, departments, branches, '
        'or divisions that need consolidated reporting.'
    )

    doc.add_heading('When to Use', level=2)
    use_cases = [
        'Parent company with subsidiaries',
        'Company with multiple departments',
        'Franchise with multiple locations',
        'Any organization needing consolidated financials'
    ]
    for case in use_cases:
        doc.add_paragraph(case, style='List Bullet')

    doc.add_heading('Setting Up Divisions', level=2)
    setup_steps = [
        'Click "Configure Divisions" to open the setup dialog',
        'Click "+ Add Division" for each entity',
        'Enter a name for each division',
        'Browse to select P&L and Balance Sheet for each division',
        'Click "Save & Continue" to return to main window'
    ]
    for i, step in enumerate(setup_steps, 1):
        doc.add_paragraph(f'{i}. {step}')

    doc.add_page_break()

    # Section 7: Update Existing Model
    doc.add_heading('7. Update Existing Model', level=1)
    doc.add_paragraph(
        'Instead of creating a new model each month, you can update an existing model with new data.'
    )

    doc.add_heading('How to Update', level=2)
    update_steps = [
        'Click Browse next to "Model File" and select your existing .xlsm file',
        'Select your NEW P&L file (containing the new month)',
        'Select your NEW Balance Sheet file (containing the new month)',
        'Click "Update Financial Model"',
        'Choose where to save (can overwrite or save as new)'
    ]
    for i, step in enumerate(update_steps, 1):
        doc.add_paragraph(f'{i}. {step}')

    doc.add_page_break()

    # Section 8: Generated Workbook
    doc.add_heading('8. Generated Excel Workbook', level=1)

    doc.add_heading('Menu Sheet', level=2)
    doc.add_paragraph(
        'The Menu sheet is your navigation hub. It contains the company name, period selector '
        '(dropdown in C7), data range display, and quick links to all sheets.'
    )

    doc.add_heading('Dashboard Sheet', level=2)
    doc.add_paragraph(
        'The Dashboard provides an executive overview with key metrics (Revenue, Gross Profit, '
        'Net Income), charts showing trends, and margin percentages.'
    )

    doc.add_heading('P&L Sheet', level=2)
    doc.add_paragraph(
        'Profit & Loss statement with monthly columns, YTD columns (Prior Year, Current Year), '
        'and variance analysis.'
    )

    doc.add_heading('Balance Sheet', level=2)
    doc.add_paragraph(
        'Balance Sheet with monthly snapshot columns, organized by Assets, Liabilities, and Equity.'
    )

    doc.add_heading('Cash Flow Statement', level=2)
    doc.add_paragraph(
        'Automatically generated from P&L and Balance Sheet data. Shows Operating, Investing, '
        'and Financing activities.'
    )

    doc.add_heading('Forecast Sheet', level=2)
    doc.add_paragraph(
        '12-month forecast with columns for Actual, Budget, Adjustment, Notes, and Forecast. '
        'Past months show Actuals, future months show Budget + Adjustments.'
    )

    doc.add_page_break()

    # Section 9: Working with Output
    doc.add_heading('9. Working with the Output', level=1)

    doc.add_heading('Changing the Viewing Period', level=2)
    doc.add_paragraph(
        'Go to the Menu sheet and click the dropdown in cell C7. Select any available month '
        'and all sheets will automatically update.'
    )

    doc.add_heading('Adding Budget Data', level=2)
    doc.add_paragraph(
        'Go to the Source_Budget sheet and enter budget amounts for each account. '
        'Columns B through M represent January through December.'
    )

    doc.add_page_break()

    # Section 10: Troubleshooting
    doc.add_heading('10. Troubleshooting', level=1)

    doc.add_heading('Common Issues', level=2)

    issues = [
        ('"The file is still open"', 'Close Excel completely and try again.'),
        ('"Could not determine date columns"', 'Ensure your Excel file has date headers in the first few rows.'),
        ('"Wrong file type detected"', 'You may have selected P&L for Balance Sheet or vice versa.'),
        ('Formulas show #REF! errors', 'Source data structure may have changed. Try regenerating.'),
        ('Charts don\'t show data', 'Do not delete row 3 (hidden helper row).')
    ]

    for issue, solution in issues:
        p = doc.add_paragraph()
        p.add_run(issue).bold = True
        p.add_run(f'\n{solution}')

    doc.add_page_break()

    # Section 11: FAQ
    doc.add_heading('11. FAQ', level=1)

    faqs = [
        ('Can I edit the generated reports?', 'Yes, but be careful with formula cells.'),
        ('How do I add a new month?', 'Use Update mode with your existing model and new data files.'),
        ('Can I use this with other accounting systems?', 'Yes, as long as exports have account names in column A and months as headers.'),
        ('Why are some sheets hidden?', 'Source sheets are hidden to keep the workbook clean. Use Settings to show/hide.'),
        ('How do I enter my budget?', 'Go to Source_Budget sheet and enter amounts.')
    ]

    for question, answer in faqs:
        p = doc.add_paragraph()
        p.add_run(f'Q: {question}').bold = True
        doc.add_paragraph(f'A: {answer}')

    # Footer
    doc.add_page_break()
    doc.add_heading('Support', level=1)
    doc.add_paragraph('Email: support@focuscfo.com')
    doc.add_paragraph('Website: www.focuscfo.com')
    doc.add_paragraph()
    p = doc.add_paragraph('CFO Financial Model Generator v3.2.13')
    p.add_run('\nCopyright 2024-2026 Focus CFO. All rights reserved.')

    # Save
    doc.save(os.path.join(os.path.dirname(__file__), 'CFO_Financial_Model_Generator_User_Manual.docx'))
    print('Created: CFO_Financial_Model_Generator_User_Manual.docx')


def create_testing_guide():
    """Create the Beta Testing Guide Word document"""
    doc = Document()

    # Title
    title = doc.add_heading('CFO Financial Model Generator', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph('Beta Testing Guide v3.2.13')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    # Introduction
    doc.add_heading('Introduction', level=1)
    doc.add_paragraph(
        'Thank you for helping test the CFO Financial Model Generator! Your feedback is essential '
        'for ensuring a quality product. This guide provides specific test scenarios, expected results, '
        'and feedback forms.'
    )
    doc.add_paragraph(
        'Important: Please complete all applicable tests and provide honest feedback. '
        'Both passes AND failures are valuable information.'
    ).bold = True

    doc.add_page_break()

    # Before You Begin
    doc.add_heading('Before You Begin', level=1)

    doc.add_heading('Test Environment Setup', level=2)
    setup = [
        'Windows PC with Excel 2016 or later installed',
        'Close all Excel windows before testing',
        'Download test files from the provided location',
        'Create a test folder (e.g., C:\\CFO_Testing\\) for generated files'
    ]
    for i, item in enumerate(setup, 1):
        doc.add_paragraph(f'{i}. {item}')

    doc.add_heading('Test Data Files Needed', level=2)
    files = [
        'Sample P&L files (single entity)',
        'Sample Balance Sheet files (single entity)',
        'Multi-division P&L files (3+ divisions)',
        'Multi-division Balance Sheet files (3+ divisions)',
        'Previous month\'s generated model (for update testing)'
    ]
    for f in files:
        p = doc.add_paragraph(f, style='List Bullet')

    doc.add_page_break()

    # Section 1: Installation Tests
    doc.add_heading('Section 1: Installation & Launch Tests', level=1)

    doc.add_heading('Test 1.1: Application Launch', level=2)
    table = doc.add_table(rows=5, cols=4)
    table.style = 'Table Grid'
    headers = table.rows[0].cells
    headers[0].text = 'Item'
    headers[1].text = 'Expected'
    headers[2].text = 'Pass/Fail'
    headers[3].text = 'Notes'

    tests = [
        ('Double-click EXE file', 'Application window opens'),
        ('Window displays correctly', 'All buttons/labels visible'),
        ('Version shows 3.2.13', 'Check bottom left corner'),
        ('No error messages', 'Clean startup')
    ]
    for i, (item, expected) in enumerate(tests, 1):
        row = table.rows[i].cells
        row[0].text = item
        row[1].text = expected

    doc.add_page_break()

    # Section 2: Single Entity Tests
    doc.add_heading('Section 2: Single Entity Mode Tests', level=1)

    doc.add_heading('Test 2.1: File Selection', level=2)
    table = doc.add_table(rows=5, cols=4)
    table.style = 'Table Grid'
    headers = table.rows[0].cells
    headers[0].text = 'Item'
    headers[1].text = 'Expected'
    headers[2].text = 'Pass/Fail'
    headers[3].text = 'Notes'

    tests = [
        ('Browse opens file dialog', 'Dialog appears'),
        ('Select P&L file', 'Path shows in text field'),
        ('File validation runs', 'No error for correct file'),
        ('Select BS file as P&L', 'Should show warning')
    ]
    for i, (item, expected) in enumerate(tests, 1):
        row = table.rows[i].cells
        row[0].text = item
        row[1].text = expected

    doc.add_heading('Test 2.2: Model Generation', level=2)
    table = doc.add_table(rows=8, cols=4)
    table.style = 'Table Grid'
    headers = table.rows[0].cells
    headers[0].text = 'Item'
    headers[1].text = 'Expected'
    headers[2].text = 'Pass/Fail'
    headers[3].text = 'Notes'

    tests = [
        ('Enter company name', 'Text accepted'),
        ('Click Generate', 'Save dialog appears'),
        ('Choose save location', 'File saves'),
        ('Progress bar updates', 'Shows progress 0-100%'),
        ('Generation completes', 'Success message'),
        ('File opens in Excel', 'Workbook opens'),
        ('Time to complete', 'Record: ___ seconds')
    ]
    for i, (item, expected) in enumerate(tests, 1):
        row = table.rows[i].cells
        row[0].text = item
        row[1].text = expected

    doc.add_heading('Test 2.3: Generated File Validation', level=2)
    table = doc.add_table(rows=9, cols=4)
    table.style = 'Table Grid'
    headers = table.rows[0].cells
    headers[0].text = 'Sheet'
    headers[1].text = 'Expected'
    headers[2].text = 'Pass/Fail'
    headers[3].text = 'Notes'

    sheets = [
        ('Menu', 'Sheet present with company name'),
        ('Dashboard', 'Sheet present with KPIs'),
        ('PL', 'Sheet present with data'),
        ('Balance_Sheet', 'Sheet present with data'),
        ('Cash_Flow', 'Sheet present'),
        ('Forecast', 'Sheet present with 12 months'),
        ('Source_PL', 'Sheet hidden'),
        ('Source_BS', 'Sheet hidden')
    ]
    for i, (sheet, expected) in enumerate(sheets, 1):
        row = table.rows[i].cells
        row[0].text = sheet
        row[1].text = expected

    doc.add_page_break()

    # Section 3: Update Tests
    doc.add_heading('Section 3: Update Existing Model Tests', level=1)

    doc.add_heading('Test 3.1: Basic Update', level=2)
    doc.add_paragraph('Use a previously generated model and new data files with one additional month.')

    table = doc.add_table(rows=8, cols=4)
    table.style = 'Table Grid'
    headers = table.rows[0].cells
    headers[0].text = 'Item'
    headers[1].text = 'Expected'
    headers[2].text = 'Pass/Fail'
    headers[3].text = 'Notes'

    tests = [
        ('Select existing model', 'File loads'),
        ('Select new P&L file', 'File validates'),
        ('Select new BS file', 'File validates'),
        ('Click Generate', 'Update starts'),
        ('Update completes', 'Success message'),
        ('New column added to PL', 'Column present'),
        ('Menu dropdown updated', 'Has new month')
    ]
    for i, (item, expected) in enumerate(tests, 1):
        row = table.rows[i].cells
        row[0].text = item
        row[1].text = expected

    doc.add_page_break()

    # Section 4: Multi-Division Tests
    doc.add_heading('Section 4: Multi-Division Mode Tests', level=1)

    doc.add_heading('Test 4.1: Division Setup', level=2)
    table = doc.add_table(rows=9, cols=4)
    table.style = 'Table Grid'
    headers = table.rows[0].cells
    headers[0].text = 'Item'
    headers[1].text = 'Expected'
    headers[2].text = 'Pass/Fail'
    headers[3].text = 'Notes'

    tests = [
        ('Click Configure Divisions', 'Dialog opens'),
        ('Add Division button', 'New row appears'),
        ('Can add 3+ divisions', 'No limit hit'),
        ('Can name each division', 'Text entry works'),
        ('Browse P&L per division', 'File dialog works'),
        ('Browse BS per division', 'File dialog works'),
        ('Remove button works', 'Row removed'),
        ('Save & Continue', 'Returns to main')
    ]
    for i, (item, expected) in enumerate(tests, 1):
        row = table.rows[i].cells
        row[0].text = item
        row[1].text = expected

    doc.add_page_break()

    # Section 5: Performance
    doc.add_heading('Section 5: Performance Tests', level=1)

    doc.add_paragraph('Record actual times for your system:')

    table = doc.add_table(rows=6, cols=4)
    table.style = 'Table Grid'
    headers = table.rows[0].cells
    headers[0].text = 'Scenario'
    headers[1].text = 'Expected'
    headers[2].text = 'Actual'
    headers[3].text = 'Pass/Fail'

    scenarios = [
        ('Single entity (3 months)', '30-90 sec'),
        ('Single entity (12 months)', '45-120 sec'),
        ('Multi-division (3 div)', '2-3 min'),
        ('Multi-division (5 div)', '3-5 min'),
        ('Update existing model', '15-45 sec')
    ]
    for i, (scenario, expected) in enumerate(scenarios, 1):
        row = table.rows[i].cells
        row[0].text = scenario
        row[1].text = expected

    doc.add_page_break()

    # Section 6: Feedback Questions
    doc.add_heading('Section 6: Feedback Questions', level=1)

    doc.add_heading('Usability', level=2)

    doc.add_paragraph('1. Was the interface intuitive?')
    for opt in ['Very intuitive', 'Somewhat intuitive', 'Confusing', 'Very confusing']:
        doc.add_paragraph(f'[ ] {opt}', style='List Bullet')

    doc.add_paragraph('2. Were error messages helpful?')
    for opt in ['Very helpful', 'Somewhat helpful', 'Not helpful', 'Did not encounter errors']:
        doc.add_paragraph(f'[ ] {opt}', style='List Bullet')

    doc.add_paragraph('3. Was the progress indicator useful?')
    for opt in ['Very useful', 'Somewhat useful', 'Not useful']:
        doc.add_paragraph(f'[ ] {opt}', style='List Bullet')

    doc.add_heading('Output Quality', level=2)

    doc.add_paragraph('4. Did the generated reports meet your expectations?')
    for opt in ['Exceeded expectations', 'Met expectations', 'Below expectations']:
        doc.add_paragraph(f'[ ] {opt}', style='List Bullet')

    doc.add_paragraph('5. Were the formulas correct?')
    for opt in ['All correct', 'Minor issues', 'Significant issues']:
        doc.add_paragraph(f'[ ] {opt}', style='List Bullet')

    doc.add_heading('Overall', level=2)

    doc.add_paragraph('6. Would you use this tool in production?')
    for opt in ['Yes, definitely', 'Yes, with improvements', 'Not sure', 'No']:
        doc.add_paragraph(f'[ ] {opt}', style='List Bullet')

    doc.add_paragraph('7. What features were most valuable?')
    doc.add_paragraph('_' * 60)

    doc.add_paragraph('8. What features were missing?')
    doc.add_paragraph('_' * 60)

    doc.add_paragraph('9. What would you change?')
    doc.add_paragraph('_' * 60)

    doc.add_page_break()

    # Bug Report Template
    doc.add_heading('Bug Report Template', level=1)
    doc.add_paragraph('If you encounter any bugs, please complete this template:')

    template = """
BUG REPORT
Date: ________________
Tester Name: ________________
Version: 3.2.13

DESCRIPTION:
What happened?
____________________________________________

STEPS TO REPRODUCE:
1. ________________________________________
2. ________________________________________
3. ________________________________________

EXPECTED BEHAVIOR:
____________________________________________

ACTUAL BEHAVIOR:
____________________________________________

ERROR MESSAGE (if any):
____________________________________________

SEVERITY:
[ ] Critical - Cannot use application
[ ] Major - Feature doesn't work
[ ] Minor - Works but not ideal
[ ] Cosmetic - Visual issue only

WORKAROUND (if found):
____________________________________________
"""
    doc.add_paragraph(template)

    doc.add_page_break()

    # Test Sign-Off
    doc.add_heading('Test Sign-Off', level=1)

    doc.add_paragraph('Tester Name: _______________________')
    doc.add_paragraph('Date Tested: _______________________')
    doc.add_paragraph('Version Tested: 3.2.13')
    doc.add_paragraph('Operating System: _______________________')
    doc.add_paragraph('Excel Version: _______________________')

    doc.add_heading('Test Summary', level=2)
    table = doc.add_table(rows=7, cols=4)
    table.style = 'Table Grid'
    headers = table.rows[0].cells
    headers[0].text = 'Section'
    headers[1].text = 'Passed'
    headers[2].text = 'Failed'
    headers[3].text = 'Skipped'

    sections = [
        '1. Installation',
        '2. Single Entity',
        '3. Update Model',
        '4. Multi-Division',
        '5. Performance',
        '6. Feedback'
    ]
    for i, section in enumerate(sections, 1):
        row = table.rows[i].cells
        row[0].text = section

    doc.add_paragraph()
    doc.add_paragraph('[ ] I have completed all applicable tests')
    doc.add_paragraph('[ ] I have documented all failures')
    doc.add_paragraph('[ ] I have provided feedback')
    doc.add_paragraph()
    doc.add_paragraph('Signature: _______________________')
    doc.add_paragraph('Date: _______________________')

    doc.add_paragraph()
    doc.add_paragraph('Please return completed test results to: support@focuscfo.com')
    doc.add_paragraph('Thank you for your valuable feedback!').bold = True

    # Save
    doc.save(os.path.join(os.path.dirname(__file__), 'Beta_Testing_Guide.docx'))
    print('Created: Beta_Testing_Guide.docx')


if __name__ == '__main__':
    create_user_manual()
    create_testing_guide()
    print('\nDone! Word documents created in docs folder.')
